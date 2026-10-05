"""ШАГ 4: Сборка Excel построчно: одна строка = один снимок товара."""
import json
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from ..core.paths import ProjectPaths
from ..core.logger import ProjectLogger


HEADERS = [
    "url", "domain", "row_num",
    "timestamp", "date", "is_duplicate",
    "input_name", "matched_name", "match_score", "match_status",
    "product_name", "sku", "brand",
    "price_current", "currency", "price_old",
    "current_price_input", "diff_percent", "diff_status",
    "availability", "confidence", "method",
    "source_file",
]


def ts_to_date(ts):
    if not ts or len(ts) < 8:
        return ""
    return f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}"


def _to_float(v):
    """Приводит значение к float."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        s = str(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
        s = "".join(c for c in s if c.isdigit() or c in ".-")
        return float(s) if s else None
    except (ValueError, TypeError):
        return None


def run(paths: ProjectPaths, logger: ProjectLogger,
        output_name: str = "report.xlsx",
        diff_threshold: float = 5.0) -> dict:
    logger.info("=" * 60)
    logger.info("ШАГ 4: Сборка Excel")
    logger.info("=" * 60)

    input_data = {}
    if paths.input_parsed.exists():
        input_data = json.loads(paths.input_parsed.read_text(encoding="utf-8"))

    input_by_url = {}
    for row in input_data.get("rows", []):
        input_by_url[row["url"]] = row

    wb = Workbook()
    ws = wb.active
    ws.title = "prices"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="FF6B00")
    header_align = Alignment(horizontal="center", vertical="center")

    ws.append(HEADERS)
    for i in range(1, len(HEADERS) + 1):
        c = ws.cell(row=1, column=i)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    total_rows = 0
    total_products = 0
    diff_high = 0
    diff_low = 0

    highlight_high = PatternFill("solid", fgColor="FFE5E5")
    highlight_low = PatternFill("solid", fgColor="E5F4FF")
    highlight_low_conf = PatternFill("solid", fgColor="FFF8E1")

    if not paths.products_dir.exists():
        logger.warn("Нет папки products/")
        wb.save(paths.exports_dir / output_name)
        return {"ok": True, "rows": 0}

    for page_dir in sorted(paths.products_dir.iterdir()):
        if not page_dir.is_dir():
            continue

        for json_file in sorted(page_dir.glob("*.json")):
            try:
                data = json.loads(json_file.read_text(encoding="utf-8"))
            except Exception as e:
                logger.warn(f"Не читается {json_file}: {e}")
                continue
            if not isinstance(data, dict):
                continue

            meta = data.get("meta", {}) or {}
            source_url = meta.get("source_url", "")
            timestamp = json_file.stem
            date = ts_to_date(timestamp)

            input_row = input_by_url.get(source_url, {})
            input_price = input_row.get("current_price")
            input_name = input_row.get("name")
            row_num = input_row.get("row_num")

            product = data.get("product", {}) or {}
            price = data.get("price", {}) or {}
            price_current = _to_float(price.get("current"))

            # Match-инфа
            match_score = meta.get("match_score")
            match_status = meta.get("match_status")
            matched_name = meta.get("matched_name")

            # Расхождение
            diff_percent = None
            diff_status = ""
            if price_current and input_price and input_price > 0:
                diff_percent = (price_current - input_price) / input_price * 100
                diff_percent = round(diff_percent, 1)
                if diff_percent > diff_threshold:
                    diff_status = "выше входной"
                elif diff_percent < -diff_threshold:
                    diff_status = "ниже входной"
                else:
                    diff_status = "совпадает"

            row = [
                source_url,
                urlparse(source_url).netloc,
                row_num,
                timestamp,
                date,
                "no",
                input_name,
                matched_name,
                match_score,
                match_status,
                product.get("name"),
                product.get("sku"),
                product.get("brand"),
                price_current,
                price.get("currency"),
                price.get("old"),
                input_price,
                diff_percent,
                diff_status,
                price.get("availability"),
                meta.get("confidence_score"),
                meta.get("method"),
                meta.get("source_file"),
            ]
            ws.append(row)
            total_rows += 1

            # Подсветка
            row_idx = ws.max_row
            fill = None
            if diff_percent is not None:
                if diff_percent > diff_threshold:
                    fill = highlight_high
                    diff_high += 1
                elif diff_percent < -diff_threshold:
                    fill = highlight_low
                    diff_low += 1
            if match_status == "low_confidence":
                fill = highlight_low_conf
            if fill:
                for col in range(1, len(HEADERS) + 1):
                    ws.cell(row=row_idx, column=col).fill = fill

            # Дубликаты
            for dup in meta.get("duplicated_dates", []) or []:
                dup_ts = dup.get("timestamp")
                dup_date = dup.get("date") or ts_to_date(dup_ts)
                dup_row = list(row)
                dup_row[3] = dup_ts
                dup_row[4] = dup_date
                dup_row[5] = "yes"
                ws.append(dup_row)
                total_rows += 1

            total_products += 1

    # Автоширина
    for i, h in enumerate(HEADERS, start=1):
        max_len = len(str(h))
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i, values_only=True):
            if row and row[0] is not None:
                max_len = max(max_len, min(len(str(row[0])), 60))
        ws.column_dimensions[get_column_letter(i)].width = max_len + 2

    ws.freeze_panes = "A2"

    out_path = paths.exports_dir / output_name
    wb.save(out_path)

    logger.success(f"Excel: {out_path}")
    logger.info(f"  Строк: {total_rows}")
    logger.info(f"  Уникальных товаров: {total_products}")
    logger.info(f"  Расхождений выше порога: {diff_high}")
    logger.info(f"  Расхождений ниже порога: {diff_low}")

    return {
        "ok": True,
        "file": str(out_path),
        "rows": total_rows,
        "products": total_products,
        "diff_high": diff_high,
        "diff_low": diff_low,
    }